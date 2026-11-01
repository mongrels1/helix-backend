"""Teacher notes page (Track LP4). Run: python -m pytest test_teacher_notes.py"""
import docx

import docx_template


def _plan():
    return {"days": [
        {"name": "Monday", "teacher_notes": {
            "where_this_fits": "Introduces unit rate with tiles.",
            "watch_for": [{"misconception": "Divides in the wrong order",
                           "looks_like": "12 ÷ 3", "teacher_response": "Which is per one?"}],
            "worked_example": {"problem": "3 : 12", "solution": "3 ÷ 12 = 0.25"},
            "practice": [{"problem": "5 : 8", "answer": "5 ÷ 8 = 0.625"}],
            "checks": ["5 ÷ 8 = 0.6 (we get 0.625)"]}},
        {"name": "Tuesday", "status": "no_school"},
    ]}


def test_notes_page_is_added_after_the_template(tmp_path):
    path = tmp_path / "plan.docx"
    d = docx.Document()
    d.add_paragraph("TEMPLATE BODY")
    d.save(path)
    assert docx_template.append_teacher_notes(str(path), _plan()) == 1
    text = [p.text for p in docx.Document(path).paragraphs if p.text.strip()]
    assert text[0] == "TEMPLATE BODY"
    assert text[1] == "Teacher notes"
    assert "Monday" in text and "Watch for" in text
    assert "Check: 5 ÷ 8 = 0.6 (we get 0.625)" in text
    assert "Tuesday" not in text  # no notes, no block


def test_no_notes_means_no_page(tmp_path):
    path = tmp_path / "plan.docx"
    docx.Document().save(path)
    before = path.read_bytes()
    assert docx_template.append_teacher_notes(str(path), {"days": [{"name": "Mon"}]}) == 0
    assert path.read_bytes() == before
