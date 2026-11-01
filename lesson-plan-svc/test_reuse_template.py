"""Last week's filled plan used as the template (Track LP7).
Run: python -m pytest test_reuse_template.py"""
import docx

import docx_template


def _plan_doc(path, filled: bool):
    d = docx.Document()
    t = d.add_table(rows=2, cols=1)
    t.rows[0].cells[0].text = "Monday"
    cell = t.rows[1].cells[0]
    cell.paragraphs[0].text = "Standard/Expectations: 7.PAR.4.1 OLD WEEK" if filled else "Standard/Expectations:"
    cell.add_paragraph("old standard detail" if filled else "")
    cell.add_paragraph("")
    cell.add_paragraph("Learning Targets:")
    cell.add_paragraph("Today, we are learning to:")
    cell.add_paragraph("Today, we are learning to compare dirty water OLD WEEK" if filled else "")
    cell.add_paragraph("Success Criteria:")
    cell.add_paragraph("I know I am successful when I can:")
    cell.add_paragraph("I can build tile models OLD WEEK" if filled else "")
    d.save(path)


def test_filled_plan_is_detected_and_cleared_back_to_the_blank(tmp_path):
    blank, filled = tmp_path / "blank.docx", tmp_path / "filled.docx"
    _plan_doc(blank, False)
    _plan_doc(filled, True)
    assert docx_template.parse_template(str(filled))["looks_filled"] is True
    assert docx_template.parse_template(str(blank))["looks_filled"] is False

    d = docx.Document(filled)
    res = docx_template.clear_previous_fill(d)
    assert res["cleared"] >= 4
    after = [p.text for p in d.tables[0].rows[1].cells[0].paragraphs]
    before = [p.text for p in docx.Document(blank).tables[0].rows[1].cells[0].paragraphs]
    assert after == before
    assert "OLD WEEK" not in " ".join(after)


def test_a_value_that_starts_like_a_label_is_cleared_not_kept(tmp_path):
    path = tmp_path / "f.docx"
    _plan_doc(path, True)
    d = docx.Document(path)
    docx_template.clear_previous_fill(d)
    texts = [p.text for p in d.tables[0].rows[1].cells[0].paragraphs]
    assert texts.count("Today, we are learning to:") == 1
    assert not any(t.startswith("Today, we are learning to compare") for t in texts)
